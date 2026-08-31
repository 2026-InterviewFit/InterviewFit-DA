import io
import os
import re
from typing import List

import fitz
from dotenv import load_dotenv
from minio import Minio
from pptx import Presentation
from pypdf import PdfReader

from app.pt_interview.schema.pt_document_schema import (
    PtDocumentAnalyzeResponse,
    PtSlideAnalyzeResult,
)
from app.pt_interview.service.pt_summary_service import (
    PtSummaryService,
)

load_dotenv()


class PtDocumentService:

    def __init__(self):

        self.minio_client = Minio(
            os.getenv(
                "MINIO_ENDPOINT",
                "localhost:9000",
            ),
            access_key=os.getenv(
                "MINIO_ACCESS_KEY"
            ),
            secret_key=os.getenv(
                "MINIO_SECRET_KEY"
            ),
            secure=os.getenv(
                "MINIO_SECURE",
                "False",
            ).lower() == "true",
        )

        self.bucket_name = os.getenv(
            "MINIO_BUCKET",
            "videos",
        )

        self.summary_service = PtSummaryService()

    def analyze_document(
        self,
        object_key: str,
    ) -> PtDocumentAnalyzeResponse:

        """
        PT 발표자료 분석

        1. MinIO 파일 다운로드
        2. 파일 형식 확인
        3. PPTX / PDF 텍스트 추출
        4. PDF 페이지 이미지 생성 및 MinIO 저장
        5. 슬라이드별 Gemini 요약
        6. 전체 슬라이드 텍스트 결합
        7. Gemini 전체 발표자료 요약
        8. 분석 결과 반환
        """

        file_bytes = self._download_from_minio(
            object_key
        )

        extension = self._get_extension(
            object_key
        )

        if extension == ".pdf":

            slides = self._extract_pdf(
                file_bytes=file_bytes,
                source_object_key=object_key,
            )

        elif extension == ".pptx":

            slides = self._extract_pptx(
                file_bytes
            )

        else:

            raise ValueError(
                "지원하지 않는 파일 형식입니다. "
                "PDF 또는 PPTX 파일만 분석할 수 있습니다."
            )

        slides = self._summarize_slides(
            slides
        )

        document_text = self._build_document_text(
            slides
        )

        document_summary = (
            self.summary_service
            .summarize_document(
                document_text
            )
        )

        return PtDocumentAnalyzeResponse(
            object_key=object_key,
            file_type=extension.replace(
                ".",
                "",
            ),
            total_slides=len(slides),
            summary=document_summary,
            slides=slides,
        )

    def _download_from_minio(
        self,
        object_key: str,
    ) -> bytes:

        response = None

        try:

            response = (
                self.minio_client
                .get_object(
                    self.bucket_name,
                    object_key,
                )
            )

            return response.read()

        except Exception as e:

            raise RuntimeError(
                f"MinIO 파일 다운로드 실패: {e}"
            ) from e

        finally:

            if response is not None:

                response.close()
                response.release_conn()

    def _get_extension(
        self,
        object_key: str,
    ) -> str:

        return os.path.splitext(
            object_key
        )[1].lower()

    def _extract_pdf(
        self,
        file_bytes: bytes,
        source_object_key: str,
    ) -> List[PtSlideAnalyzeResult]:

        """
        PDF 처리

        - pypdf: 텍스트 추출
        - PyMuPDF: 각 페이지 PNG 렌더링
        - 생성된 PNG는 MinIO 저장
        """

        reader = PdfReader(
            io.BytesIO(
                file_bytes
            )
        )

        pdf_document = fitz.open(
            stream=file_bytes,
            filetype="pdf",
        )

        slides = []

        try:

            for index, page in enumerate(
                reader.pages,
                start=1,
            ):

                raw_text = (
                    page.extract_text()
                    or ""
                )

                extracted_text = (
                    self._clean_extracted_text(
                        raw_text
                    )
                )

                image_object_key = (
                    self._render_and_upload_pdf_page(
                        pdf_document=pdf_document,
                        page_index=index - 1,
                        source_object_key=source_object_key,
                        slide_number=index,
                    )
                )

                slides.append(
                    PtSlideAnalyzeResult(
                        slide_number=index,
                        title=self._extract_title(
                            extracted_text
                        ),
                        extracted_text=extracted_text,
                        summary=None,
                        image_object_key=image_object_key,
                    )
                )

        finally:

            pdf_document.close()

        return slides

    def _render_and_upload_pdf_page(
        self,
        pdf_document,
        page_index: int,
        source_object_key: str,
        slide_number: int,
    ) -> str:

        """
        PDF 한 페이지를 PNG 이미지로 변환 후
        MinIO에 저장한다.
        """

        page = pdf_document.load_page(
            page_index
        )

        # 2배 정도의 해상도로 렌더링
        matrix = fitz.Matrix(
            2.0,
            2.0,
        )

        pixmap = page.get_pixmap(
            matrix=matrix,
            alpha=False,
        )

        image_bytes = pixmap.tobytes(
            "png"
        )

        base_name = os.path.splitext(
            source_object_key
        )[0]

        image_object_key = (
            f"{base_name}/slides/"
            f"slide-{slide_number}.png"
        )

        self.minio_client.put_object(
            bucket_name=self.bucket_name,
            object_name=image_object_key,
            data=io.BytesIO(
                image_bytes
            ),
            length=len(
                image_bytes
            ),
            content_type="image/png",
        )

        return image_object_key

    def _extract_pptx(
        self,
        file_bytes: bytes,
    ) -> List[PtSlideAnalyzeResult]:

        presentation = Presentation(
            io.BytesIO(
                file_bytes
            )
        )

        slides = []

        for index, slide in enumerate(
            presentation.slides,
            start=1,
        ):

            text_parts = []

            for shape in slide.shapes:

                if not hasattr(
                    shape,
                    "text",
                ):
                    continue

                text = shape.text.strip()

                if text:

                    text_parts.append(
                        text
                    )

            extracted_text = (
                self._clean_extracted_text(
                    "\n".join(
                        text_parts
                    )
                )
            )

            slides.append(
                PtSlideAnalyzeResult(
                    slide_number=index,
                    title=self._extract_title(
                        extracted_text
                    ),
                    extracted_text=extracted_text,
                    summary=None,

                    # PPTX 이미지는 현재 생성하지 않음
                    image_object_key=None,
                )
            )

        return slides

    def _summarize_slides(
        self,
        slides: List[
            PtSlideAnalyzeResult
        ],
    ) -> List[
        PtSlideAnalyzeResult
    ]:

        """
        슬라이드별 Gemini 요약 생성
        """

        summarized_slides = []

        for slide in slides:

            summary = (
                self.summary_service
                .summarize_slide(
                    slide_number=slide.slide_number,
                    title=slide.title,
                    extracted_text=slide.extracted_text,
                )
            )

            summarized_slides.append(
                PtSlideAnalyzeResult(
                    slide_number=slide.slide_number,
                    title=slide.title,
                    extracted_text=slide.extracted_text,
                    summary=summary,
                    image_object_key=(
                        slide.image_object_key
                    ),
                )
            )

        return summarized_slides

    def _clean_extracted_text(
        self,
        text: str,
    ) -> str:

        """
        PDF/PPTX에서 추출된 텍스트 정리.

        - 불필요한 공백 제거
        - 여러 빈 줄 제거
        - 줄 앞뒤 공백 제거
        """

        if not text:
            return ""

        text = text.replace(
            "\r\n",
            "\n",
        )

        text = text.replace(
            "\r",
            "\n",
        )

        lines = []

        for line in text.splitlines():

            line = re.sub(
                r"[ \t]+",
                " ",
                line,
            ).strip()

            if not line:
                continue

            lines.append(
                line
            )

        return "\n".join(
            lines
        ).strip()

    def _extract_title(
        self,
        text: str,
    ) -> str | None:

        """
        제목 후보 추출.

        01, 02 같은 숫자만 있는 줄은 제목에서 제외한다.
        """

        if not text:
            return None

        lines = [
            line.strip()
            for line in text.splitlines()
            if line.strip()
        ]

        for line in lines:

            # 01 / 02 / 1 / 2 같은 값 제외
            if re.fullmatch(
                r"\d{1,3}",
                line,
            ):
                continue

            # 너무 짧은 의미 없는 텍스트 제외
            if len(line) <= 1:
                continue

            return line[:200]

        return None

    def _build_document_text(
        self,
        slides: List[
            PtSlideAnalyzeResult
        ],
    ) -> str:

        """
        Gemini 전체 발표자료 요약용 텍스트.
        """

        text_parts = []

        for slide in slides:

            if not slide.extracted_text:
                continue

            text_parts.append(
                f"""
[슬라이드 {slide.slide_number}]

제목:
{slide.title or "없음"}

원본 내용:
{slide.extracted_text}

슬라이드 핵심 요약:
{slide.summary or "없음"}
""".strip()
            )

        return "\n\n".join(
            text_parts
        )